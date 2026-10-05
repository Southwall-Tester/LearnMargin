"""Bounded local document extraction with explicit opt-in for Office conversion."""
from __future__ import annotations

import io
import os
import posixpath
import re
import shutil
import signal
import stat
import struct
import subprocess
import tempfile
import warnings
import zipfile
import zlib
from pathlib import Path
from urllib.parse import unquote, urlsplit
from xml.etree import ElementTree
from xml.parsers import expat

from bs4 import BeautifulSoup
from PIL import Image, ImageOps

from .models import Document, SourceUnit

MAX_UPLOAD_BYTES = 50 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 250 * 1024 * 1024
MAX_MEMBER_BYTES = 50 * 1024 * 1024
MAX_ZIP_ENTRIES = 5000
MAX_UNITS = 300
MAX_IMAGE_PIXELS = 40_000_000
MAX_IMAGES = 600
MAX_TEXT_CHARS = 2_000_000
CHUNK_CHARS = 12_000
CONVERSION_TIMEOUT = 120
ZIP_READ_CHUNK = 64 * 1024
SUPPORTED_EXTENSIONS = {
    ".pdf", ".pptx", ".docx", ".epub", ".txt", ".md", ".markdown",
    ".html", ".htm", ".csv", ".tsv", ".rst", ".json", ".log",
    ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff", ".tif", ".gif",
    ".ppt", ".doc", ".odt", ".odp", ".rtf",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff", ".tif", ".gif"}
LEGACY_EXTENSIONS = {".ppt", ".doc", ".odt", ".odp", ".rtf"}


class IngestionError(ValueError):
    """A user-actionable document error, safe to display in the local UI."""


def find_libreoffice() -> str | None:
    """Find an existing installation; never install Office or start a GUI."""
    for name in ("libreoffice", "soffice", "soffice.exe"):
        if binary := shutil.which(name):
            return binary
    for root in (os.environ.get("ProgramFiles"), os.environ.get("ProgramFiles(x86)")):
        if root:
            candidate = Path(root) / "LibreOffice" / "program" / "soffice.exe"
            if candidate.is_file():
                return str(candidate)
    return None


def local_office_enabled() -> bool:
    """This opt-in is local configuration, never an upload/request parameter."""
    return os.environ.get("LEARNMARGIN_ALLOW_LOCAL_OFFICE", "").strip() == "1"


def _xml_guard():
    """Reject declarations in the XML parser, after its encoding detection."""
    parser = expat.ParserCreate()

    def reject(*_args):
        raise IngestionError("文档包含不支持的 XML 外部声明，请重新导出。")

    parser.StartDoctypeDeclHandler = reject
    parser.EntityDeclHandler = reject
    parser.ExternalEntityRefHandler = reject
    parser.SetParamEntityParsing(expat.XML_PARAM_ENTITY_PARSING_NEVER)
    return parser


def _zip_chunks(raw, item: zipfile.ZipInfo, remaining_budget: int):
    """Validate the real compressed stream, not ZipInfo's attacker-controlled size.

    ZipExtFile truncates to file_size, and an unbounded read can inflate far more
    before that truncation. Validate independently with zlib's output limit; do
    not call flush(), whose output is unbounded.
    """
    raw.seek(item.header_offset)
    header = raw.read(30)
    if len(header) != 30 or header[:4] != b"PK\x03\x04":
        raise IngestionError("文档压缩包的文件头损坏，请重新导出。")
    fields = struct.unpack("<4s5H3I2H", header)
    if fields[3] != item.compress_type or fields[2] != item.flag_bits:
        raise IngestionError("文档压缩包的文件头与目录不一致，请重新导出。")
    raw.seek(fields[-2] + fields[-1], os.SEEK_CUR)
    compressed_left = item.compress_size
    decoder = zlib.decompressobj(-zlib.MAX_WBITS) if item.compress_type == zipfile.ZIP_DEFLATED else None
    pending = b""
    actual = 0
    checksum = 0
    while compressed_left or pending:
        if not pending:
            pending = raw.read(min(ZIP_READ_CHUNK, compressed_left))
            if not pending:
                raise IngestionError("文档压缩包内容不完整，请重新导出。")
            compressed_left -= len(pending)
        # The extra byte detects false size declarations without first expanding
        # the whole member, even when its forged CRC matches the truncated data.
        limit = min(ZIP_READ_CHUNK, item.file_size - actual + 1,
                    MAX_MEMBER_BYTES - actual + 1, remaining_budget - actual + 1)
        if decoder:
            chunk = decoder.decompress(pending, max(1, limit))
            pending = decoder.unconsumed_tail
        else:
            chunk, pending = pending[:max(1, limit)], pending[max(1, limit):]
        actual += len(chunk)
        if actual > MAX_MEMBER_BYTES or actual > remaining_budget:
            raise IngestionError("文档实际解压内容超过限制，请拆分材料再上传。")
        if actual > item.file_size:
            raise IngestionError("文档压缩包声明大小与实际内容不一致，请重新导出。")
        checksum = zlib.crc32(chunk, checksum)
        yield chunk
        if decoder and decoder.eof:
            if decoder.unused_data or pending or compressed_left:
                raise IngestionError("文档压缩包包含多余的压缩数据，请重新导出。")
            break
    if (decoder is not None and not decoder.eof) or actual != item.file_size or checksum != item.CRC:
        raise IngestionError("文档压缩包的实际内容与校验信息不一致，请重新导出。")


def _check_zip(path: Path) -> None:
    """Bound actual expansion and validate XML before using document libraries."""
    with zipfile.ZipFile(path) as archive, path.open("rb") as raw:
        entries = archive.infolist()
        if len(entries) > MAX_ZIP_ENTRIES:
            raise IngestionError(f"压缩包条目过多，最多允许 {MAX_ZIP_ENTRIES} 项。")
        if sum(item.file_size for item in entries) > MAX_UNCOMPRESSED_BYTES:
            raise IngestionError("文档解压后超过 250 MB，请拆分材料再上传。")
        names: set[str] = set()
        offsets: set[int] = set()
        xml_parts: set[str] = set()
        xml_extensions: set[str] = set()

        def content_type(tag, attributes):
            if not attributes.get("ContentType", "").lower().endswith(("+xml", "/xml")):
                return
            if tag.rsplit(":", 1)[-1] == "Override":
                xml_parts.add(unquote(attributes.get("PartName", "")).lstrip("/"))
            elif tag.rsplit(":", 1)[-1] == "Default":
                xml_extensions.add(attributes.get("Extension", "").lower())
            if len(xml_parts) + len(xml_extensions) > MAX_ZIP_ENTRIES:
                raise IngestionError("文档包含过多的内容类型声明，请拆分材料再上传。")

        actual_total = 0
        # Read the guarded OOXML manifest first, regardless of ZIP entry order.
        # Relationships can assign XML content types to arbitrary suffixes.
        for item in sorted(entries, key=lambda entry: entry.filename != "[Content_Types].xml"):
            name = item.filename.replace("\\", "/")
            if (name.startswith("/") or ":" in name or ".." in name.split("/")
                    or name in names or item.orig_filename != item.filename
                    or stat.S_ISLNK(item.external_attr >> 16) or item.header_offset in offsets):
                raise IngestionError("文档压缩包包含不安全或重复的文件路径。")
            names.add(name)
            offsets.add(item.header_offset)
            if item.flag_bits & 1:
                raise IngestionError("无法读取加密的文档压缩包，请先保存为未加密副本。")
            if item.file_size > MAX_MEMBER_BYTES:
                raise IngestionError("文档内单个文件超过 50 MB，请缩小图片或拆分材料。")
            if item.file_size > 1024 * 1024 and item.file_size > max(1, item.compress_size) * 1000:
                raise IngestionError("文档压缩比例异常，请重新导出材料。")
            if item.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}:
                raise IngestionError("文档使用不支持的 ZIP 压缩方式，请用原应用重新导出。")
            if name.lower().endswith("vbaproject.bin"):
                raise IngestionError("暂不导入包含宏的 Office 文档，请另存为不含宏的副本。")
            # open() validates names, extra fields and overlapping compressed
            # ranges but deliberately never reads/decompresses member data here.
            with archive.open(item):
                pass
            is_xml = (name.lower().endswith((".xml", ".rels", ".opf", ".ncx")) or name in xml_parts
                      or name.rsplit(".", 1)[-1].lower() in xml_extensions)
            xml = _xml_guard() if is_xml else None
            if name == "[Content_Types].xml":
                xml.StartElementHandler = content_type
            for chunk in _zip_chunks(raw, item, MAX_UNCOMPRESSED_BYTES - actual_total):
                actual_total += len(chunk)
                if xml is not None:
                    xml.Parse(chunk, False)
            if xml is not None:
                xml.Parse(b"", True)


class _Collector:
    def __init__(self, output: Path):
        self.output = output
        self.units: list[SourceUnit] = []
        self.warnings: list[str] = []
        self.image_count = 0
        self.text_count = 0

    def warn(self, message: str) -> None:
        if message not in self.warnings:
            self.warnings.append(message)

    def add(self, label: str, text: str, images: list[str] | None = None) -> None:
        if len(self.units) >= MAX_UNITS:
            raise IngestionError(f"材料超过 {MAX_UNITS} 个内容单元，请先拆分再上传。")
        self.text_count += len(text)
        if self.text_count > MAX_TEXT_CHARS:
            raise IngestionError("材料文字量超过 200 万字符，请按学习范围拆分后上传。")
        self.units.append(SourceUnit(index=len(self.units) + 1, label=label, text=text.strip(),
                                     image_paths=images or []))

    def image(self, source: bytes | Image.Image) -> str:
        if self.image_count >= MAX_IMAGES:
            raise IngestionError(f"材料图片超过 {MAX_IMAGES} 张，请拆分后上传。")
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            original = Image.open(io.BytesIO(source)) if isinstance(source, bytes) else source
            picture = None
            try:
                if original.width * original.height > MAX_IMAGE_PIXELS:
                    raise IngestionError("单张图片超过 4000 万像素，请缩小图片后上传。")
                picture = ImageOps.exif_transpose(original)
                picture.thumbnail((1800, 1800))
                if picture.mode not in ("RGB", "RGBA"):
                    picture = picture.convert("RGBA" if "transparency" in picture.info else "RGB")
                self.image_count += 1
                name = f"image-{self.image_count:04d}.png"
                picture.save(self.output / name, "PNG")
                return name
            finally:
                if picture is not None:
                    picture.close()
                if isinstance(source, bytes):
                    original.close()

    def text_chunks(self, text: str, label: str, images: list[str] | None = None) -> None:
        # Bound model inputs without silently dropping long chapters or inventing page numbers.
        chunks: list[str] = []
        remaining = text.strip()
        while len(remaining) > CHUNK_CHARS:
            cut = remaining.rfind("\n", CHUNK_CHARS // 2, CHUNK_CHARS)
            cut = cut if cut > 0 else CHUNK_CHARS
            chunks.append(remaining[:cut])
            remaining = remaining[cut:].lstrip()
        if remaining or not chunks:
            chunks.append(remaining)
        for index, chunk in enumerate(chunks):
            suffix = f"（续 {index + 1}）" if index else ""
            self.add(f"{label}{suffix}", chunk, images if index == 0 else None)
        if len(chunks) > 1:
            self.warn("较长章节已分成连续内容段；这些段号不是原文件页码。")


def _pdf(path: Path, result: _Collector) -> None:
    import pypdfium2 as pdfium
    from pypdf import PdfReader
    from pypdf.errors import LimitReachedError
    from pypdf.generic import ArrayObject

    reader = PdfReader(path)
    if reader.is_encrypted and not reader.decrypt(""):
        raise IngestionError("PDF 需要密码，请先导出不加密的副本。")
    count = len(reader.pages)
    if count > MAX_UNITS:
        raise IngestionError(f"PDF 超过 {MAX_UNITS} 页，请先导出所需页码范围。")
    if count == 0:
        raise IngestionError("PDF 没有可读取的页面。")
    with pdfium.PdfDocument(path) as rendered:
        for index, page in enumerate(reader.pages):
            try:
                # Large content streams can have disproportionate parser memory costs.
                content = page.get("/Contents")
                if content is not None:
                    content = content.get_object()
                    streams = content if isinstance(content, ArrayObject) else [content]
                    decoded_size = 0
                    for stream in streams:
                        decoded_size += len(stream.get_object().get_data())
                        if decoded_size > MAX_MEMBER_BYTES:
                            raise IngestionError(f"PDF 第 {index + 1} 页内容流过大，请重新导出此页。")
                text = page.extract_text() or ""
            except IngestionError:
                raise
            except LimitReachedError as exc:
                raise IngestionError(f"PDF 第 {index + 1} 页解压内容过大，请重新导出此页。") from exc
            except Exception:
                text = ""
                result.warn("部分 PDF 页面无法抽取文字，已保留页面图像。")
            rendered_page = rendered[index]
            try:
                width, height = rendered_page.get_size()
                if min(width, height) <= 0:
                    raise IngestionError(f"PDF 第 {index + 1} 页尺寸无效。")
                bitmap = rendered_page.render(scale=min(2.0, 1800 / max(width, height)))
                try:
                    image = result.image(bitmap.to_pil())
                finally:
                    bitmap.close()
            finally:
                rendered_page.close()
            result.add(f"第 {index + 1} 页", text, [image])
            if not text.strip():
                result.warn("部分 PDF 页面没有可提取文字，可能是扫描件；需要支持图像的模型理解这些页面。")
    result.warn("页码按 PDF 文件实际页序计算，可能不同于页面上印刷的编号；公式、图表请结合页面图像核对。")


def _pptx(path: Path, result: _Collector) -> None:
    from pptx import Presentation

    presentation = Presentation(path)
    if len(presentation.slides) > MAX_UNITS:
        raise IngestionError(f"演示文稿超过 {MAX_UNITS} 张幻灯片，请拆分后上传。")

    def collect_shapes(shapes, texts: list[str], images: list[str]) -> None:
        for shape in shapes:
            if hasattr(shape, "shapes"):
                collect_shapes(shape.shapes, texts, images)
            if getattr(shape, "has_text_frame", False):
                texts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False):
                texts.append("\n".join(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows))
            if hasattr(shape, "image"):
                try:
                    images.append(result.image(shape.image.blob))
                except (OSError, ValueError) as exc:
                    if isinstance(exc, IngestionError):
                        raise
                    result.warn("部分 Office 图片格式无法转换，请将含这些图的幻灯片另存为 PDF 再上传。")
            if getattr(shape, "has_chart", False):
                chart = shape.chart
                if chart.has_title and chart.chart_title.has_text_frame:
                    texts.append("图表标题：" + chart.chart_title.text_frame.text)
                result.warn("PPTX 图表、公式、SmartArt、形状布局和母版不保证完整提取；需要准确视觉理解时请另存为 PDF。")

    for index, slide in enumerate(presentation.slides, 1):
        texts: list[str] = []
        images: list[str] = []
        collect_shapes(slide.shapes, texts, images)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                texts.append("讲者备注：\n" + notes.text)
        result.add(f"第 {index} 张幻灯片", "\n\n".join(texts), images)
    result.warn("PPTX 按幻灯片提取文字、表格、内嵌图片和讲者备注；不保留完整布局、公式、母版及动画。建议将视觉信息丰富的材料导出为 PDF。")


def _docx(path: Path, result: _Collector) -> None:
    from docx import Document as WordDocument
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = WordDocument(path)
    title = "开头内容"
    texts: list[str] = []
    images: list[str] = []

    def flush() -> None:
        if texts or images:
            result.text_chunks("\n\n".join(texts), title, images.copy())
            texts.clear()
            images.clear()

    for element in document.element.body:
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, document)
            style_name = paragraph.style.name if paragraph.style else ""
            outline = element.find(".//" + qn("w:outlineLvl"))
            if (re.match(r"^(Heading|标题)\s*\d", style_name, re.I) or outline is not None) and paragraph.text.strip():
                flush()
                title = paragraph.text.strip()
            if paragraph.text.strip():
                texts.append(paragraph.text)
        elif element.tag == qn("w:tbl"):
            table = Table(element, document)
            texts.append("\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows))
        for blip in element.iter(qn("a:blip")):
            relationship = blip.get(qn("r:embed"))
            if relationship and relationship in document.part.related_parts:
                try:
                    images.append(result.image(document.part.related_parts[relationship].blob))
                except (OSError, ValueError) as exc:
                    if isinstance(exc, IngestionError):
                        raise
                    result.warn("部分 Word 图片格式无法转换，请导出为 PDF 后上传。")
    flush()
    result.warn("DOCX 按标题和连续内容段划分，不是真实印刷页码；页眉页脚、浮动文本框、修订与公式可能不完整。需要按页学习时请上传导出的 PDF。")


def _clean_html(content: bytes | str) -> BeautifulSoup:
    soup = BeautifulSoup(content, "html.parser")
    for node in soup(["script", "style", "iframe", "object", "embed", "nav", "noscript"]):
        node.decompose()
    return soup


def _safe_xml(content: bytes):
    parser = _xml_guard()
    parser.Parse(content, True)
    return ElementTree.fromstring(content)


def _epub(path: Path, result: _Collector) -> None:
    namespace = {"c": "urn:oasis:names:tc:opendocument:xmlns:container",
                 "o": "http://www.idpf.org/2007/opf"}
    with zipfile.ZipFile(path) as archive:
        if "META-INF/encryption.xml" in archive.namelist():
            result.warn("电子书包含加密声明，部分字体或内容可能无法读取；不尝试绕过 DRM。")
        container = _safe_xml(archive.read("META-INF/container.xml"))
        rootfile = container.find(".//c:rootfile", namespace)
        if rootfile is None:
            raise IngestionError("EPUB 缺少内容入口，请重新导出电子书。")
        package_path = rootfile.attrib["full-path"]
        package = _safe_xml(archive.read(package_path))
        base = posixpath.dirname(package_path)
        manifest = {item.attrib["id"]: item for item in package.findall("o:manifest/o:item", namespace)}
        spine = package.findall("o:spine/o:itemref", namespace)
        if len(spine) > MAX_UNITS:
            raise IngestionError(f"电子书超过 {MAX_UNITS} 个章节文件，请先导出需要的章节。")
        for index, item in enumerate(spine, 1):
            entry = manifest.get(item.attrib.get("idref", ""))
            if entry is None:
                raise IngestionError("EPUB 的章节索引不完整，请重新导出。")
            href = entry.attrib.get("href", "")
            if urlsplit(href).scheme or urlsplit(href).netloc:
                result.warn("已忽略电子书中的外部章节链接。")
                continue
            member = posixpath.normpath(posixpath.join(base, unquote(urlsplit(href).path)))
            if member.startswith("../") or member.startswith("/"):
                raise IngestionError("电子书章节路径无效。")
            soup = _clean_html(archive.read(member))
            heading = soup.find(["h1", "h2", "h3"])
            title = heading.get_text(" ", strip=True) if heading else f"第 {index} 章段"
            images: list[str] = []
            for element in soup.find_all("img"):
                href = element.get("src", "")
                parts = urlsplit(href)
                if parts.scheme or parts.netloc:
                    result.warn("电子书中的远程或内联图片未加载；只读取包内图片。")
                    continue
                image_path = posixpath.normpath(posixpath.join(posixpath.dirname(member), unquote(parts.path)))
                if image_path.startswith(("../", "/")):
                    continue
                try:
                    images.append(result.image(archive.read(image_path)))
                except (KeyError, OSError, ValueError) as exc:
                    if isinstance(exc, IngestionError):
                        raise
                    result.warn("电子书中部分图片无法读取。")
            result.text_chunks(soup.get_text("\n", strip=True), title, images)
    result.warn("EPUB 按书内阅读顺序划分章段；章段号不是阅读器动态排版后的页码。")


def _text(path: Path, kind: str, result: _Collector) -> None:
    data = path.read_bytes()
    try:
        text = data.decode("utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = data.decode("gb18030")
            result.warn("文本按 GB18030 编码读取，请检查预览中的字符是否正确。")
        except UnicodeDecodeError as exc:
            raise IngestionError("无法识别文本编码，请将文件另存为 UTF-8 后上传。") from exc
    if "\x00" in text:
        raise IngestionError("文件含有二进制数据，请确认格式或另存为 UTF-8 文本。")
    if kind in (".html", ".htm"):
        soup = _clean_html(text)
        if soup.find("img"):
            result.warn("HTML 仅提取文字，不读取外部或本地引用图片；需要保留视觉信息时请先打印为 PDF。")
        text = soup.get_text("\n", strip=True)
    result.text_chunks(text, "文本内容")
    result.warn("文本文件按连续内容段划分，段号不是页码。")


def _run_converter(command: list[str]) -> int:
    """Bound the conversion process tree, including LibreOffice's child process."""
    in_worker = os.environ.get("LEARNMARGIN_EXTRACTION_WORKER") == "1"
    options = ({"creationflags": subprocess.CREATE_NO_WINDOW} if os.name == "nt"
               else {"start_new_session": not in_worker})
    with subprocess.Popen(command, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, **options) as process:
        try:
            return process.wait(timeout=CONVERSION_TIMEOUT)
        except subprocess.TimeoutExpired as exc:
            if os.name == "nt":
                # The PID belongs to the subprocess just created here, never a user's Office session.
                subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"],
                               capture_output=True, check=False, timeout=10,
                               creationflags=subprocess.CREATE_NO_WINDOW)
            elif not in_worker:
                try:
                    os.killpg(process.pid, signal.SIGKILL)
                except ProcessLookupError:
                    pass
            process.kill()
            process.wait(timeout=10)
            raise IngestionError("LibreOffice 转换超过 120 秒，请自行导出 PDF 或拆分材料。") from exc


def _legacy(path: Path, kind: str, result: _Collector) -> None:
    if not local_office_enabled():
        raise IngestionError("旧格式转换默认关闭，请先导出 PDF 后导入；可信文件可由本机配置启用。")
    executable = find_libreoffice()
    if not executable:
        raise IngestionError(f"读取 {kind} 需要 LibreOffice。请安装 LibreOffice，或先用 Office 导出为 PDF/PPTX/DOCX 再上传。")
    if kind in {".odt", ".odp"}:
        _check_zip(path)
    with tempfile.TemporaryDirectory(prefix="learnmargin-convert-") as directory:
        work = Path(directory)
        profile = work / "profile"
        profile.mkdir()
        # A dedicated profile avoids reusing a user's Office session; level 3 disables macros.
        (profile / "user").mkdir()
        (profile / "user" / "registrymodifications.xcu").write_text(
            '<?xml version="1.0" encoding="UTF-8"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
            '<item oor:path="/org.openoffice.Office.Common/Security/Scripting">'
            '<prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop>'
            '</item><item oor:path="/org.openoffice.Office.Common/Load">'
            '<prop oor:name="Update" oor:op="fuse"><value>0</value></prop></item></oor:items>',
            encoding="utf-8",
        )
        input_file = work / f"source{kind}"
        shutil.copyfile(path, input_file)
        command = [executable, f"-env:UserInstallation={profile.as_uri()}", "--headless",
                   "--nologo", "--nodefault", "--nolockcheck", "--norestore",
                   "--convert-to", "pdf", "--outdir", str(work), str(input_file)]
        try:
            return_code = _run_converter(command)
        except OSError as exc:
            raise IngestionError("无法启动 LibreOffice，请检查安装或自行导出为 PDF。") from exc
        converted = work / "source.pdf"
        if return_code != 0 or not converted.is_file():
            raise IngestionError("LibreOffice 未能转换材料，请用原应用导出为 PDF 后上传。")
        if converted.stat().st_size > MAX_UPLOAD_BYTES:
            raise IngestionError("转换后的 PDF 超过 50 MB，请拆分材料。")
        _pdf(converted, result)
        result.warn("此文件已通过本机 LibreOffice 转为 PDF；页码指转换结果，字体与分页可能不同于原应用。")


def extract_document(path: Path, output_dir: Path, document_id: str,
                     original_name: str | None = None) -> Document:
    """Read a local upload and return honest source units plus relative PNG paths."""
    path = Path(path)
    output_dir = Path(output_dir)
    name = Path(original_name or path.name).name
    kind = Path(name).suffix.lower()
    if not path.is_file():
        raise IngestionError("找不到上传文件，请重新上传。")
    if path.stat().st_size > MAX_UPLOAD_BYTES:
        raise IngestionError("文件超过 50 MB，请压缩图片或按学习范围拆分后上传。")
    if path.stat().st_size == 0:
        raise IngestionError("文件为空，请检查后重新上传。")
    if kind not in SUPPORTED_EXTENSIONS:
        raise IngestionError(f"暂不支持 {kind or '无扩展名'} 文件，请转为 PDF、PPTX、DOCX、EPUB、文本或图片。")
    output_dir.mkdir(parents=True, exist_ok=True)
    result = _Collector(output_dir)
    try:
        if kind in {".pptx", ".docx", ".epub"}:
            _check_zip(path)
        if kind == ".pdf":
            _pdf(path, result)
            unit_label = "页"
        elif kind == ".pptx":
            _pptx(path, result)
            unit_label = "幻灯片"
        elif kind == ".docx":
            _docx(path, result)
            unit_label = "节段"
        elif kind == ".epub":
            _epub(path, result)
            unit_label = "章段"
        elif kind in IMAGE_EXTENSIONS:
            with Image.open(path) as source:
                frames = getattr(source, "n_frames", 1)
                if frames > MAX_UNITS:
                    raise IngestionError(f"图片包含超过 {MAX_UNITS} 帧，请拆分后上传。")
                # Multipage TIFF is commonly a scanned handout; GIF represents one illustration.
                count = frames if kind in {".tif", ".tiff"} else 1
                for index in range(count):
                    source.seek(index)
                    result.add(f"第 {index + 1} 张图片", "", [result.image(source)])
                if frames > 1 and count == 1:
                    result.warn("动态图像仅保留第一帧；如需其他帧请分别上传。")
            result.warn("图片未执行 OCR；需要支持图像的模型理解其内容。")
            unit_label = "图片"
        elif kind in LEGACY_EXTENSIONS:
            _legacy(path, kind, result)
            unit_label = "转换后页"
        else:
            _text(path, kind, result)
            unit_label = "内容段"
        if not result.units or not any(unit.text or unit.image_paths for unit in result.units):
            raise IngestionError("没有提取到可学习的文字或图片，请检查材料或导出为 PDF。")
        return Document(id=document_id, name=name, kind=kind.lstrip("."), unit_label=unit_label,
                        units=result.units, warnings=result.warnings)
    except IngestionError:
        raise
    except (Image.DecompressionBombWarning, Image.DecompressionBombError) as exc:
        raise IngestionError("图片像素数过大，请缩小图片后上传。") from exc
    except Exception as exc:
        raise IngestionError("无法读取文件，可能已损坏、加密或扩展名与内容不符。请用原应用重新导出后上传。") from exc
