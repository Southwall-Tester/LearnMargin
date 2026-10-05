from __future__ import annotations

import io
import stat
import struct
import subprocess
import zipfile
import zlib
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document as WordDocument
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from learnmargin import ingestion
from learnmargin.ingestion import IngestionError, extract_document


def image_bytes() -> bytes:
    stream = io.BytesIO()
    with Image.new("RGB", (80, 60), "navy") as picture:
        picture.save(stream, "PNG")
    return stream.getvalue()


def make_pdf(path: Path, *, pages: int = 2, password: str | None = None) -> None:
    writer = PdfWriter()
    for index in range(pages):
        page = writer.add_blank_page(width=300, height=400)
        if index == 0:
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                     NameObject("/Subtype"): NameObject("/Type1"),
                                     NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
            stream = DecodedStreamObject()
            stream.set_data(b"BT /F1 18 Tf 30 320 Td (Conditional probability) Tj ET")
            page[NameObject("/Contents")] = writer._add_object(stream)
    if password:
        writer.encrypt(password)
    writer.write(path)


def assert_images_are_local(document, output: Path) -> None:
    for unit in document.units:
        for name in unit.image_paths:
            assert Path(name).name == name
            assert (output / name).is_file()
            with Image.open(output / name) as picture:
                assert max(picture.size) <= 1800


def test_pdf_preserves_text_page_order_and_image_for_every_page(tmp_path):
    source = tmp_path / "input.pdf"
    output = tmp_path / "extracted"
    make_pdf(source)
    document = extract_document(source, output, "sample")
    assert document.unit_label == "页"
    assert [unit.index for unit in document.units] == [1, 2]
    assert "Conditional probability" in document.units[0].text
    assert document.units[1].text == ""
    assert all(len(unit.image_paths) == 1 for unit in document.units)
    assert any("扫描件" in warning for warning in document.warnings)
    assert_images_are_local(document, output)


def test_pdf_password_and_page_limits_are_actionable(tmp_path, monkeypatch):
    source = tmp_path / "locked.pdf"
    make_pdf(source, password="private")
    with pytest.raises(IngestionError, match="密码"):
        extract_document(source, tmp_path / "out", "locked")
    make_pdf(source)
    monkeypatch.setattr(ingestion, "MAX_UNITS", 1)
    with pytest.raises(IngestionError, match="超过 1 页"):
        extract_document(source, tmp_path / "out", "long")


def test_pptx_preserves_slides_tables_notes_and_embedded_images(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text_frame.text = "条件概率"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(3), Inches(1)).table
    table.cell(0, 0).text = "事件"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "12"
    slide.shapes.add_picture(io.BytesIO(image_bytes()), Inches(5), Inches(1))
    slide.notes_slide.notes_text_frame.text = "先解释分母。"
    presentation.slides.add_slide(presentation.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "练习"
    source = tmp_path / "source.pptx"
    presentation.save(source)
    output = tmp_path / "out"
    document = extract_document(source, output, "slides")
    assert document.unit_label == "幻灯片"
    assert len(document.units) == 2
    assert "条件概率" in document.units[0].text
    assert "A | 12" in document.units[0].text
    assert "先解释分母" in document.units[0].text
    assert document.units[1].text == "练习"
    assert_images_are_local(document, output)


def test_docx_uses_headings_and_preserves_table_and_picture(tmp_path):
    document = WordDocument()
    document.add_heading("条件概率", level=1)
    document.add_paragraph("已知条件改变样本空间。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A∩B"
    table.cell(0, 1).text = "6"
    document.add_picture(io.BytesIO(image_bytes()))
    document.add_heading("独立性", level=1)
    document.add_paragraph("比较交集概率与乘积。")
    source = tmp_path / "source.docx"
    document.save(source)
    output = tmp_path / "out"
    result = extract_document(source, output, "word")
    assert result.unit_label == "节段"
    assert [unit.label for unit in result.units] == ["条件概率", "独立性"]
    assert "A∩B | 6" in result.units[0].text
    assert len(result.units[0].image_paths) == 1
    assert any("不是真实印刷页码" in warning for warning in result.warnings)
    assert_images_are_local(result, output)


def make_epub(path: Path) -> None:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", "application/epub+zip")
        archive.writestr("META-INF/container.xml", '''<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
          <rootfiles><rootfile full-path="OPS/book.opf"/></rootfiles></container>''')
        archive.writestr("OPS/book.opf", '''<package xmlns="http://www.idpf.org/2007/opf"><manifest>
          <item id="b" href="second.xhtml" media-type="application/xhtml+xml"/>
          <item id="a" href="first.xhtml" media-type="application/xhtml+xml"/>
          </manifest><spine><itemref idref="a"/><itemref idref="b"/></spine></package>''')
        archive.writestr("OPS/first.xhtml", '<html><body><h1>第一章</h1><p>先学条件。</p>'
                          '<img src="media/picture.png"/><img src="https://example.com/track.png"/>'
                          '<script>NEVER EXECUTE</script></body></html>')
        archive.writestr("OPS/second.xhtml", "<html><body><h1>第二章</h1><p>再学独立。</p></body></html>")
        archive.writestr("OPS/media/picture.png", image_bytes())


def test_epub_obeys_spine_not_zip_order_and_ignores_remote_resources(tmp_path):
    source = tmp_path / "book.epub"
    output = tmp_path / "out"
    make_epub(source)
    result = extract_document(source, output, "epub")
    assert [unit.label for unit in result.units] == ["第一章", "第二章"]
    assert "NEVER EXECUTE" not in result.units[0].text
    assert len(result.units[0].image_paths) == 1
    assert any("远程" in warning for warning in result.warnings)
    assert_images_are_local(result, output)


@pytest.mark.parametrize("member,payload,error", [
    ("../outside.xml", "x", "路径"),
    ("word/document.xml", '<!DOCTYPE doc [<!ENTITY a "xx">]><doc>&a;</doc>', "XML"),
    ("word/vbaProject.bin", "fake macro", "宏"),
])
def test_rejects_unsafe_office_packages_before_parser(tmp_path, member, payload, error):
    source = tmp_path / "unsafe.docx"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr(member, payload)
    with pytest.raises(IngestionError, match=error):
        extract_document(source, tmp_path / "out", "bad")
    assert not (tmp_path / "outside.xml").exists()


def test_zip_expansion_limit_precedes_parser(tmp_path, monkeypatch):
    source = tmp_path / "huge.epub"
    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("large.txt", "A" * 101)
    monkeypatch.setattr(ingestion, "MAX_UNCOMPRESSED_BYTES", 100)
    with pytest.raises(IngestionError, match="解压"):
        extract_document(source, tmp_path / "out", "bad")


def test_text_encoding_chunking_and_original_upload_name(tmp_path, monkeypatch):
    source = tmp_path / "upload.data"
    source.write_bytes("这是第一段。\n这是第二段。\n第三段。".encode("gb18030"))
    monkeypatch.setattr(ingestion, "CHUNK_CHARS", 10)
    result = extract_document(source, tmp_path / "out", "text", "笔记.md")
    assert result.name == "笔记.md"
    assert result.kind == "md"
    assert len(result.units) >= 2
    assert "第三段" in "".join(unit.text for unit in result.units)
    assert any("GB18030" in warning for warning in result.warnings)


def test_html_drops_executable_and_remote_content(tmp_path):
    source = tmp_path / "source.html"
    source.write_text('<html><style>HIDE ME</style><script>RUN ME</script>'
                      '<h1>主题</h1><p>正文</p><img src="file:///private/file.png">'
                      '<iframe src="https://example.com"></iframe></html>', encoding="utf-8")
    result = extract_document(source, tmp_path / "out", "html")
    assert result.units[0].text == "主题\n正文"
    assert not result.units[0].image_paths


def test_images_are_bounded_and_tiff_preserves_frames(tmp_path, monkeypatch):
    source = tmp_path / "scan.tiff"
    with Image.new("RGB", (40, 30), "white") as first, Image.new("RGB", (40, 30), "red") as second:
        first.save(source, save_all=True, append_images=[second])
    result = extract_document(source, tmp_path / "out", "scan")
    assert len(result.units) == 2
    assert all(unit.text == "" for unit in result.units)
    assert_images_are_local(result, tmp_path / "out")
    monkeypatch.setattr(ingestion, "MAX_IMAGE_PIXELS", 100)
    with pytest.raises(IngestionError, match="像素"):
        extract_document(source, tmp_path / "out2", "big")


def test_legacy_missing_dependency_and_timeout_are_actionable(tmp_path, monkeypatch):
    monkeypatch.setenv("LEARNMARGIN_ALLOW_LOCAL_OFFICE", "1")
    source = tmp_path / "old.doc"
    source.write_bytes(b"old document fixture")
    monkeypatch.setattr(ingestion, "find_libreoffice", lambda: None)
    with pytest.raises(IngestionError, match="需要 LibreOffice"):
        extract_document(source, tmp_path / "out", "old")
    monkeypatch.setattr(ingestion, "find_libreoffice", lambda: "soffice")

    def timeout(command):
        assert "--headless" in command
        raise IngestionError("LibreOffice 转换超过 120 秒，请自行导出 PDF 或拆分材料。")

    monkeypatch.setattr(ingestion, "_run_converter", timeout)
    with pytest.raises(IngestionError, match="120 秒"):
        extract_document(source, tmp_path / "out", "old")


def test_converter_timeout_terminates_only_its_own_process_tree(monkeypatch):
    calls = []

    class Process:
        pid = 123456

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def wait(self, timeout):
            if timeout == 120:
                raise subprocess.TimeoutExpired("soffice", timeout)
            return -1

        def kill(self):
            calls.append("kill")

    monkeypatch.setattr(ingestion.subprocess, "Popen", lambda *args, **kwargs: Process())
    if ingestion.os.name == "nt":
        monkeypatch.setattr(ingestion.subprocess, "run", lambda command, **kwargs: calls.append(command))
    else:
        monkeypatch.setattr(ingestion.os, "killpg", lambda pid, sig: calls.append(["group", pid]))
    with pytest.raises(IngestionError, match="120 秒"):
        ingestion._run_converter(["soffice", "--headless"])
    assert "kill" in calls
    assert any("123456" in str(call) for call in calls)


@pytest.mark.parametrize("name,content,message", [
    ("empty.txt", b"", "为空"),
    ("source.exe", b"binary", "暂不支持"),
    ("corrupt.docx", b"broken", "无法读取"),
    ("binary.txt", b"hello\x00world", "二进制"),
])
def test_bad_inputs_produce_chinese_errors_without_paths(tmp_path, name, content, message):
    source = tmp_path / name
    source.write_bytes(content)
    with pytest.raises(IngestionError, match=message) as caught:
        extract_document(source, tmp_path / "out", "bad")
    assert str(tmp_path) not in str(caught.value)


def test_upload_size_limit(tmp_path, monkeypatch):
    source = tmp_path / "large.txt"
    source.write_text("a" * 100, encoding="utf-8")
    monkeypatch.setattr(ingestion, "MAX_UPLOAD_BYTES", 99)
    with pytest.raises(IngestionError, match="超过 50 MB"):
        extract_document(source, tmp_path / "out", "large")


def test_forged_zip_size_and_matching_truncated_crc_are_rejected_before_document_parser(tmp_path, monkeypatch):
    source = tmp_path / "forged.docx"
    # 8 MB expands from a few KB; both directory and local header lie about the
    # uncompressed length. A CRC of the first 100 bytes also fools chunked
    # ZipExtFile reads, which truncate output to the declared size.
    payload = b"A" * (8 * 1024 * 1024)
    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", payload)
    data = bytearray(source.read_bytes())
    central = data.index(b"PK\x01\x02")
    checksum = zlib.crc32(payload[:100])
    struct.pack_into("<I", data, 14, checksum)
    struct.pack_into("<I", data, 22, 100)
    struct.pack_into("<I", data, central + 16, checksum)
    struct.pack_into("<I", data, central + 24, 100)
    source.write_bytes(data)
    monkeypatch.setattr(ingestion, "_docx", lambda *_: pytest.fail("unsafe archive reached document parser"))
    with pytest.raises(IngestionError, match="声明大小与实际内容不一致"):
        extract_document(source, tmp_path / "out", "bad")


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16", "utf-16-be"])
@pytest.mark.parametrize("member", ["word/document.xml", "word/content.bin"])
def test_encoded_xml_declarations_are_rejected_before_office_parser(tmp_path, monkeypatch, encoding, member):
    source = tmp_path / "entity.docx"
    payload = ('<?xml version="1.0" encoding="UTF-16"?>' if encoding.startswith("utf-16") else
               '<?xml version="1.0" encoding="UTF-8"?>')
    payload += '<!DOCTYPE doc [<!ENTITY value "expanded">]><doc>&value;</doc>'
    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(member, payload.encode(encoding))
        archive.writestr("[Content_Types].xml",
                          '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                          f'<Override PartName="/{member}" ContentType="application/example+xml"/>'
                          '</Types>')
    monkeypatch.setattr(ingestion, "_docx", lambda *_: pytest.fail("DTD reached document parser"))
    with pytest.raises(IngestionError, match="XML"):
        extract_document(source, tmp_path / "out", "bad")


def test_safe_xml_rejects_encoded_external_entities_without_resolving_them():
    payload = ('<?xml version="1.0" encoding="UTF-16"?>'
               '<!DOCTYPE doc [<!ENTITY value SYSTEM "file:///synthetic-test-only.txt">]>'
               '<doc>&value;</doc>').encode("utf-16")
    with pytest.raises(IngestionError, match="XML"):
        ingestion._safe_xml(payload)
    assert ingestion._safe_xml('<?xml version="1.0" encoding="UTF-16"?><doc>正文</doc>'.encode("utf-16")).text == "正文"


def test_zip_streaming_descriptors_zip64_and_utf16_xml_still_work(tmp_path):
    class Unseekable(io.BytesIO):
        def seek(self, *args):
            raise io.UnsupportedOperation("not seekable")

    stream = Unseekable()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        with archive.open("word/document.xml", "w", force_zip64=True) as member:
            member.write('<?xml version="1.0" encoding="UTF-16"?><doc>正文</doc>'.encode("utf-16"))
        archive.writestr("word/media/image.png", image_bytes())
    source = tmp_path / "streamed.docx"
    source.write_bytes(stream.getvalue())
    ingestion._check_zip(source)


@pytest.mark.parametrize("compression", [zipfile.ZIP_BZIP2, zipfile.ZIP_LZMA])
def test_non_office_zip_compression_is_rejected(tmp_path, compression):
    source = tmp_path / "compressed.docx"
    with zipfile.ZipFile(source, "w", compression) as archive:
        archive.writestr("word/document.xml", "<doc/>")
    with pytest.raises(IngestionError, match="压缩方式"):
        ingestion._check_zip(source)


def test_zip_symbolic_link_is_not_treated_as_a_regular_member(tmp_path):
    source = tmp_path / "linked.docx"
    link = zipfile.ZipInfo("word/document.xml")
    link.create_system = 3
    link.external_attr = (stat.S_IFLNK | 0o777) << 16
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr(link, "../../outside.xml")
    with pytest.raises(IngestionError, match="路径"):
        ingestion._check_zip(source)


@pytest.mark.parametrize("value", [None, "0", "true"])
def test_legacy_conversion_requires_explicit_local_opt_in(tmp_path, monkeypatch, value):
    if value is None:
        monkeypatch.delenv("LEARNMARGIN_ALLOW_LOCAL_OFFICE", raising=False)
    else:
        monkeypatch.setenv("LEARNMARGIN_ALLOW_LOCAL_OFFICE", value)
    source = tmp_path / "old.doc"
    source.write_bytes(b"synthetic old document")
    monkeypatch.setattr(ingestion, "find_libreoffice", lambda: pytest.fail("converter should not be consulted"))
    monkeypatch.setattr(ingestion, "_run_converter", lambda *_: pytest.fail("converter must not be started"))
    with pytest.raises(IngestionError, match="默认关闭"):
        extract_document(source, tmp_path / "out", "old")


@pytest.mark.parametrize("in_worker", [False, True])
def test_unix_converter_stays_in_extraction_workers_process_group(monkeypatch, in_worker):
    calls = []

    class Process:
        pid = 123456

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def wait(self, timeout):
            if timeout == ingestion.CONVERSION_TIMEOUT:
                raise subprocess.TimeoutExpired("soffice", timeout)
            return -1

        def kill(self):
            calls.append("kill")

    def start(command, **kwargs):
        assert kwargs["start_new_session"] is not in_worker
        return Process()

    # Mock only ingestion's OS view, so this Unix process-lifecycle regression
    # also runs on Windows without changing pathlib/pytest's global platform.
    monkeypatch.setattr(ingestion, "os", SimpleNamespace(
        name="posix", environ={"LEARNMARGIN_EXTRACTION_WORKER": "1" if in_worker else "0"},
        killpg=lambda pid, signal: calls.append("kill group")))
    monkeypatch.setattr(ingestion, "signal", SimpleNamespace(SIGKILL=9))
    monkeypatch.setattr(ingestion.subprocess, "Popen", start)
    with pytest.raises(IngestionError, match="120 秒"):
        ingestion._run_converter(["soffice", "--headless"])
    assert ("kill group" in calls) is not in_worker
    assert "kill" in calls
